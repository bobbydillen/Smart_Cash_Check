"""report.py - builds the multi-sheet Excel reconciliation workbook (live formulas)."""
from __future__ import annotations
from collections import defaultdict
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reconcile_lib import ANOMALY_TOLERANCE, SPLIT_SEP, portions

A = "Arial"
HFILL = PatternFill("solid", fgColor="1F4E78")
HFONT = Font(name=A, bold=True, color="FFFFFF", size=10)
TITLE = Font(name=A, bold=True, size=14)
SUB = Font(name=A, bold=True, size=12, color="1F4E78")
NORM = Font(name=A, size=10)
BOLD = Font(name=A, size=10, bold=True)
ITAL = Font(name=A, italic=True, size=9)
ITAL_RED = Font(name=A, italic=True, size=9, color="C00000")
FLAG = PatternFill("solid", fgColor="FFF2CC")
BAD = PatternFill("solid", fgColor="F8CBAD")
GOOD = PatternFill("solid", fgColor="E2EFDA")
_t = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_t, right=_t, top=_t, bottom=_t)
MONEY = "#,##0.00"


def _header(ws, row, labels):
    for i, h in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font, c.border = HFILL, HFONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def _hhmm(t):
    return t.strftime("%H:%M") if isinstance(t, datetime) else ""


def source_shop_majority(results_by_shop):
    """Per payment source: which shop received most of its matched money."""
    tot = defaultdict(lambda: defaultdict(float))
    for shop, res in results_by_shop.items():
        for m in res:
            for src, amt in portions(m):
                tot[src][shop] += amt
    return {s: max(d, key=d.get) for s, d in tot.items()}


def is_info(status):
    """Matched, but worth knowing (cross-shop / combined / split) - not a problem to chase."""
    return status.startswith("Matched - ") and "Review" not in status


def classify(m, majority):
    """-> (status, confidence, note, fill)"""
    b = m.bill
    if m.payment is None and b.card <= 0:
        return ("Matched", "High", "Cash sale - no digital payment recorded, nothing to verify.", None)
    if m.payment is None:
        return ("Unmatched - Review", "Low",
                "No matching digital payment found in any source. Possibly paid cash but keyed as card, "
                "or settled through a source not uploaded.", BAD)
    excess = b.card - b.net
    if excess > ANOMALY_TOLERANCE:
        return ("Matched - Excess Digital (Review)", "Low",
                f"Digital receipt confirmed (Rs.{b.card:,.0f}) but exceeds the Net bill by Rs.{excess:,.0f}. "
                "Check for a typing error or an under-billed sale; compare with the printed bill.", BAD)
    if m.combo == "combined":
        return ("Matched - Combined Payment", "Medium", m.combo_note, GOOD)
    if m.combo == "split":
        return ("Matched - Split Across Payments", "Medium", m.combo_note, GOOD)
    home = majority.get(m.payment.source)
    if home and home != b.shop:
        return ("Matched - Cross-shop Source", "High",
                f"Paid via {m.payment.source}, which mostly serves {home}. Money confirmed received; "
                f"reassign to {b.shop} in shop-wise books.", GOOD)
    if b.net - b.card > 0:
        return ("Matched", "High", "Split payment - cash portion accepted, not investigated.", None)
    return ("Matched", "High", "", None)


def build_workbook(date_label, shop_names, results_by_shop, unused_pool, source_records, out_path):
    majority = source_shop_majority(results_by_shop)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    refs = {s: _bills_sheet(wb, s, results_by_shop[s], majority) for s in shop_names}
    src_refs = {n: _source_sheet(wb, n, r) for n, r in source_records.items()}
    _discrepancy_sheet(wb, shop_names, results_by_shop, unused_pool, majority)
    _summary_sheet(wb, date_label, shop_names, results_by_shop, refs, source_records, src_refs, unused_pool)

    order = ["A. Summary", "E-F. Discrepancies"] + [r["name"] for r in refs.values()] + [r["name"] for r in src_refs.values()]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(out_path)


def _bills_sheet(wb, shop, results, majority):
    name = f"{shop} - Bills"[:31]
    ws = wb.create_sheet(name)
    ws["A1"] = f"{shop.upper()} - bill-by-bill match"
    ws["A1"].font = TITLE
    ws["A2"] = ("Split payments are not flagged: only the recorded Card/Digital amount is checked against "
                "actual receipts; the cash remainder is accepted at face value.")
    ws["A2"].font = ITAL
    heads = ["Bill No", "Bill Time", "Till/System", "Net Bill Amt", "Recorded Card/Digital", "Implied Cash (Net-Card)",
             "Matched Payment", "Payment Source", "Reference", "Diff (Match-Card)", "Status", "Confidence", "Notes"]
    _header(ws, 4, heads)
    first = 5
    for m in sorted(results, key=lambda x: x.bill.time):
        b = m.bill
        status, conf, note, fill = classify(m, majority)
        ws.append([b.billno, _hhmm(b.time), b.system, b.net, b.card, b.net - b.card,
                   m.payment.amount if m.payment else None, m.payment.source if m.payment else None,
                   m.payment.ref if m.payment else None, m.amount_diff, status, conf, note])
        r = ws.max_row
        for c in range(1, 14):
            cell = ws.cell(row=r, column=c)
            cell.border, cell.font = BORDER, NORM
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill:
                cell.fill = fill
        for c in (4, 5, 6, 7, 10):
            ws.cell(row=r, column=c).number_format = MONEY
    last = ws.max_row
    tot = last + 1
    ws.cell(row=tot, column=1, value="TOTAL").font = BOLD
    for c, L in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
        cell = ws.cell(row=tot, column=c, value=f"=SUM({L}{first}:{L}{last})")
        cell.font, cell.number_format = BOLD, MONEY
    _widths(ws, [11, 9, 14, 12, 14, 14, 13, 24, 22, 12, 30, 11, 70])
    ws.freeze_panes = "A5"
    return dict(name=name, first=first, last=last, total=tot)


def _source_sheet(wb, source, records):
    name = source[:31]
    ws = wb.create_sheet(name)
    ws["A1"] = f"{source.upper()} - raw records"
    ws["A1"].font = TITLE
    ws["A2"] = f"{len(records)} records. Sources without a per-row time (bank statements) show the statement date."
    ws["A2"].font = ITAL
    _header(ws, 4, ["Time / Date", "Amount", "Reference", "Shop Hint"])
    for p in sorted(records, key=lambda x: (0, x.order_key.timestamp()) if isinstance(x.order_key, datetime) else (1, float(x.order_key or 0))):
        t = p.order_key.strftime("%H:%M:%S") if isinstance(p.order_key, datetime) else p.time_str
        ws.append([t, p.amount, p.ref, p.shop_hint or ""])
        r = ws.max_row
        ws.cell(row=r, column=2).number_format = MONEY
        for c in range(1, 5):
            ws.cell(row=r, column=c).border, ws.cell(row=r, column=c).font = BORDER, NORM
    last = ws.max_row
    tot = last + 1
    ws.cell(row=tot, column=1, value="TOTAL").font = BOLD
    c = ws.cell(row=tot, column=2, value=f"=SUM(B5:B{last})" if last >= 5 else 0)
    c.font, c.number_format = BOLD, MONEY
    _widths(ws, [14, 14, 30, 18])
    ws.freeze_panes = "A5"
    return dict(name=name, first=5, last=last, total=tot)


def _discrepancy_sheet(wb, shop_names, results_by_shop, unused_pool, majority):
    ws = wb.create_sheet("E-F. Discrepancies")
    ws["A1"] = "DISCREPANCY INVESTIGATION"
    ws["A1"].font = TITLE
    notes = ["Split payments are NOT treated as discrepancies - only the digital portion is verified.",
             "Every payment source is matched against ALL shops' bills together, so cross-shop routing "
             "(a sale from one shop paid through another shop's QR/machine) is detected automatically."]
    for i, n in enumerate(notes):
        ws.cell(row=3 + i, column=1, value=n).font = ITAL
    ws.cell(row=6, column=1, value="ITEMS NEEDING A LOOK (largest first)").font = SUB
    heads = ["Shop", "Bill No.", "Bill Time", "Net Bill", "Recorded Digital", "Matched Payment",
             "Payment Source", "Status", "Confidence", "Most likely explanation / what to check"]
    _header(ws, 7, heads)
    rows = []
    for shop in shop_names:
        for m in results_by_shop[shop]:
            status, conf, note, fill = classify(m, majority)
            if status == "Matched":
                continue
            size = (m.bill.card - m.bill.net) if "Excess" in status else m.bill.card
            rows.append((1 if is_info(status) else 0, -abs(size), shop, m, status, conf, note, fill))
    rows.sort(key=lambda x: (x[0], x[1]))
    for _, _, shop, m, status, conf, note, fill in rows:
        ws.append([shop, m.bill.billno, _hhmm(m.bill.time), m.bill.net, m.bill.card,
                   m.payment.amount if m.payment else None, m.payment.source if m.payment else None,
                   status, conf, note])
        r = ws.max_row
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.border, cell.font, cell.fill = BORDER, NORM, fill or PatternFill()
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c in (4, 5, 6):
            ws.cell(row=r, column=c).number_format = MONEY
    if not rows:
        ws.append(["None - every bill matched cleanly."])
    r = ws.max_row + 3
    ws.cell(row=r, column=1, value="UNASSIGNED DIGITAL RECEIPTS").font = SUB
    ws.cell(row=r + 1, column=1, value="Payments left after matching every bill from every shop: money that arrived "
            "but fits no bill. Check for a missing bill, a duplicate payment, or a tip.").font = ITAL
    if unused_pool:
        _header(ws, r + 2, ["Amount", "Source", "Reference", "Time/Date"])
        for p in unused_pool:
            ws.append([p.amount, p.source, p.ref, p.time_str])
            rr = ws.max_row
            ws.cell(row=rr, column=1).number_format = MONEY
            for c in range(1, 5):
                ws.cell(row=rr, column=c).fill, ws.cell(row=rr, column=c).border = FLAG, BORDER
    else:
        ws.cell(row=r + 2, column=1, value="None - every digital payment matched to a bill.").font = NORM
    _widths(ws, [16, 12, 10, 12, 14, 14, 24, 30, 11, 80])
    ws.freeze_panes = "A8"


def _summary_sheet(wb, date_label, shops, results_by_shop, refs, source_records, src_refs, unused_pool):
    ws = wb.create_sheet("A. Summary")
    ws["A1"] = f"SALES & PAYMENT RECONCILIATION - {date_label}"
    ws["A1"].font = TITLE
    ws["A2"] = "Shops: " + ", ".join(shops)
    ws["A2"].font = ITAL
    ws["A3"] = ("Split payments are not flagged - only the digital portion is checked. Payment sources are not "
                "pre-assigned to shops; cross-shop routing is detected automatically from the matching.")
    ws["A3"].font = ITAL_RED
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[3].height = 28

    r = 5
    ws.cell(row=r, column=1, value="A. OVERALL SUMMARY").font = SUB
    r += 1
    _header(ws, r, ["Metric", "Amount (Rs.)"])
    r += 1
    net_rows = []
    for s in shops:
        ref = refs[s]
        ws.cell(row=r, column=1, value=f"{s} - total net sales")
        ws.cell(row=r, column=2, value=f"='{ref['name']}'!D{ref['total']}")
        net_rows.append(r)
        r += 1
    net_total = r
    ws.cell(row=r, column=1, value="COMBINED TOTAL NET SALES").font = BOLD
    ws.cell(row=r, column=2, value="=" + "+".join(f"B{i}" for i in net_rows)).font = BOLD
    r += 1
    src_rows = []
    for n, ref in src_refs.items():
        ws.cell(row=r, column=1, value=f"{n} - total received")
        ws.cell(row=r, column=2, value=f"='{ref['name']}'!B{ref['total']}")
        src_rows.append(r)
        r += 1
    dig_total = r
    ws.cell(row=r, column=1, value="COMBINED DIGITAL RECEIPTS (all sources)").font = BOLD
    ws.cell(row=r, column=2, value="=" + "+".join(f"B{i}" for i in src_rows)).font = BOLD
    r += 1
    un_row = r
    ws.cell(row=r, column=1, value="  less: unassigned (matches no bill)")
    ws.cell(row=r, column=2, value=sum(p.amount for p in unused_pool))
    r += 1
    cash_row = r
    ws.cell(row=r, column=1, value="IMPLIED CASH (net sales - digital matched to bills)").font = BOLD
    ws.cell(row=r, column=2, value=f"=B{net_total}-(B{dig_total}-B{un_row})").font = BOLD
    for rr in range(net_rows[0], r + 1):
        ws.cell(row=rr, column=2).number_format = MONEY
        ws.cell(row=rr, column=1).border = ws.cell(row=rr, column=2).border = BORDER
    for rr, f in ((net_total, GOOD), (dig_total, GOOD), (cash_row, FLAG)):
        ws.cell(row=rr, column=1).fill = ws.cell(row=rr, column=2).fill = f
    ws.cell(row=r + 1, column=1, value="Implied cash is an arithmetic residual, NOT a counted cash figure (no cash report supplied). "
            "It includes all normal split-payment cash portions.").font = ITAL_RED

    r += 4
    ws.cell(row=r, column=1, value="B. SHOP-WISE").font = SUB
    r += 1
    _header(ws, r, ["Metric"] + shops)
    r += 1
    def row(label, fn, money=True):
        nonlocal r
        ws.cell(row=r, column=1, value=label).border = BORDER
        for i, s in enumerate(shops, 2):
            c = ws.cell(row=r, column=i, value=fn(s, refs[s]))
            c.border = BORDER
            if money:
                c.number_format = MONEY
        r += 1
    rng = lambda ref, col: f"'{ref['name']}'!{col}{ref['first']}:{col}{ref['last']}"
    row("Bills", lambda s, ref: f"=COUNTA({rng(ref,'A')})", False)
    row("Total net sales", lambda s, ref: f"='{ref['name']}'!D{ref['total']}")
    row("Recorded card/digital (till)", lambda s, ref: f"='{ref['name']}'!E{ref['total']}")
    row("Digital confirmed (matched to this shop's bills)", lambda s, ref: f"='{ref['name']}'!G{ref['total']}")
    row("Implied cash (net - confirmed digital)", lambda s, ref: f"='{ref['name']}'!D{ref['total']}-'{ref['name']}'!G{ref['total']}")
    row("Bills with no matching payment", lambda s, ref: f'=COUNTIF({rng(ref,"K")},"Unmatched*")', False)
    row("Bills with digital above bill amount (review)", lambda s, ref: f'=COUNTIF({rng(ref,"K")},"*Excess*")', False)
    row("Bills paid via another shop's source", lambda s, ref: f'=COUNTIF({rng(ref,"K")},"*Cross-shop*")', False)
    row("Bills paid in a combined payment (with other bills)", lambda s, ref: f'=COUNTIF({rng(ref,"K")},"*Combined*")', False)
    row("Bills paid in several parts (split across payments)", lambda s, ref: f'=COUNTIF({rng(ref,"K")},"*Split Across*")', False)

    r += 2
    ws.cell(row=r, column=1, value="C. PAYMENT SOURCES - how much of each source paid for each shop's bills").font = SUB
    r += 1
    _header(ws, r, ["Source", "Transactions"] + [f"-> {s}" for s in shops] + ["Source total"])
    r += 1
    first_src = r
    for n, recs in source_records.items():
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=len(recs))
        for i, s in enumerate(shops, 3):
            ref = refs[s]
            ws.cell(row=r, column=i, value=f'=SUMIFS({rng(ref,"G")},{rng(ref,"H")},$A{r})').number_format = MONEY
        ws.cell(row=r, column=3 + len(shops), value=f"='{src_refs[n]['name']}'!B{src_refs[n]['total']}").number_format = MONEY
        for c in range(1, 4 + len(shops)):
            ws.cell(row=r, column=c).border = BORDER
        r += 1
    ws.cell(row=r, column=1, value="Bills paid through several different sources (split)")
    for i, s in enumerate(shops, 3):
        ref = refs[s]
        ws.cell(row=r, column=i, value=f'=SUMIFS({rng(ref,"G")},{rng(ref,"H")},"*{SPLIT_SEP}*")').number_format = MONEY
    for c in range(1, 4 + len(shops)):
        ws.cell(row=r, column=c).border = BORDER
    r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = BOLD
    for c in range(2, 4 + len(shops)):
        L = get_column_letter(c)
        cell = ws.cell(row=r, column=c, value=f"=SUM({L}{first_src}:{L}{r-1})")
        cell.font = BOLD
        if c > 2:
            cell.number_format = MONEY
    ws.cell(row=r + 1, column=1, value="Off-diagonal amounts (e.g. a bank account titled for one shop paying another shop's bills) "
            "are cross-shop routing.").font = ITAL
    _widths(ws, [58, 18, 20, 20, 20, 20])
    ws.freeze_panes = "A5"
